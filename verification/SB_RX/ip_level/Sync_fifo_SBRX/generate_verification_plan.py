# Verification Plan Generator for Sync_FIFO_SBRX
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_verification_plan():
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles definition
    HEADER_FILL = PatternFill(start_color="1B4F72", end_color="1B4F72", fill_type="solid") # Dark Navy Blue from user screenshot
    HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    CAT_FILL = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid") # Dark Charcoal/Blue for Category rows
    CAT_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    
    SUB_HEADER_FILL = PatternFill(start_color="2980B9", end_color="2980B9", fill_type="solid")
    SUB_HEADER_FONT = Font(name="Calibri", size=10, bold=True, color="FFFFFF")

    ALT_ROW_FILL = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")
    WHITE_ROW_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    
    DIRECT_FILL = PatternFill(start_color="E8F8F5", end_color="E8F8F5", fill_type="solid") # Light mint green
    RANDOM_FILL = PatternFill(start_color="FEF9E7", end_color="FEF9E7", fill_type="solid") # Light amber
    
    REGULAR_FONT = Font(name="Calibri", size=10, color="000000")
    BOLD_FONT = Font(name="Calibri", size=10, bold=True, color="000000")
    CODE_FONT = Font(name="Consolas", size=9, color="1B4F72", bold=True)
    
    DASH_TITLE_FONT = Font(name="Calibri", size=16, bold=True, color="1B4F72")
    SECTION_TITLE_FONT = Font(name="Calibri", size=12, bold=True, color="1B4F72")
    
    THIN_BORDER_SIDE = Side(border_style="thin", color="D5D8DC")
    THIN_BORDER = Border(left=THIN_BORDER_SIDE, right=THIN_BORDER_SIDE, top=THIN_BORDER_SIDE, bottom=THIN_BORDER_SIDE)
    
    HEADER_BORDER_SIDE = Side(border_style="medium", color="154360")
    HEADER_BORDER = Border(left=THIN_BORDER_SIDE, right=THIN_BORDER_SIDE, top=HEADER_BORDER_SIDE, bottom=HEADER_BORDER_SIDE)

    ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ALIGN_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ALIGN_RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=True)

    # Test cases data organized by Category
    categories = [
        {
            "id": "CAT_RST",
            "name": "Reset & Initialization Verification",
            "desc": "Verification of asynchronous active-low reset assertion, deassertion, recovery, and pre-reset state clearing under idle and dynamic traffic conditions.",
            "tests": [
                {
                    "tc_id": "TC_RST_001",
                    "name": "Power-On Asynchronous Reset Assertion & Deassertion",
                    "type": "direct",
                    "desc": "Apply active-low reset (sys_rst=0) at start of simulation prior to active clock edges. Hold for 5 clock cycles, then deassert sys_rst=1 synchronously to lclk.",
                    "exp": "DUT initializes cleanly: empty=1, full=0, rd_data=0. Write/read pointers reset to address 0. Internal occupancy equals 0."
                },
                {
                    "tc_id": "TC_RST_002",
                    "name": "Asynchronous Reset Assertion While Idle",
                    "type": "direct",
                    "desc": "With FIFO already in idle empty state, assert sys_rst=0 asynchronously midway between rising clock edges.",
                    "exp": "empty=1 and full=0 remain asserted/deasserted without glitching. rd_data is driven to 0 immediately (asynchronous reset effect)."
                },
                {
                    "tc_id": "TC_RST_003",
                    "name": "Asynchronous Reset with Partially Occupied FIFO",
                    "type": "direct",
                    "desc": "Write 16 valid data words into FIFO (occupancy=16, empty=0, full=0). Assert sys_rst=0 asynchronously. Release reset. Execute read attempts.",
                    "exp": "All 16 stored items are cleared asynchronously. empty asserts to 1, full=0, rd_data resets to 0. Post-reset reads are rejected; pre-reset data is never returned."
                },
                {
                    "tc_id": "TC_RST_004",
                    "name": "Asynchronous Reset at Full Capacity",
                    "type": "direct",
                    "desc": "Fill FIFO to maximum capacity of 64 entries (full=1, empty=0). Assert sys_rst=0 asynchronously between clock edges.",
                    "exp": "full deasserts immediately to 0; empty asserts to 1; rd_data resets to 0. Internal storage is completely flushed."
                },
                {
                    "tc_id": "TC_RST_005",
                    "name": "Asynchronous Reset During Active Write Request",
                    "type": "direct",
                    "desc": "Assert wr_en=1 with test data on rising clock edge. Drive sys_rst=0 asynchronously within the same cycle before the next clock edge.",
                    "exp": "Write operation is aborted. After reset recovery, FIFO remains completely empty (empty=1); no partial or corrupted word is stored."
                },
                {
                    "tc_id": "TC_RST_006",
                    "name": "Asynchronous Reset During Active Read Request",
                    "type": "direct",
                    "desc": "Pre-load 8 entries. Drive rd_en=1 to trigger an active read. Concurrently assert sys_rst=0 asynchronously.",
                    "exp": "Read operation is aborted. rd_data is forced to 0 (reset state) rather than capturing the popped item. FIFO initializes to empty."
                },
                {
                    "tc_id": "TC_RST_007",
                    "name": "Narrow Asynchronous Reset Pulse Verification (Glitch / Min Pulse Width)",
                    "type": "direct",
                    "desc": "Apply a narrow active-low reset pulse (sys_rst=0) of duration less than half a clock period between rising clock edges.",
                    "exp": "Asynchronous reset successfully resets internal FIFO state, pointers, rd_data=0, and flags (empty=1, full=0). Normal operation resumes on next clock edge."
                },
                {
                    "tc_id": "TC_RST_008",
                    "name": "Post-Reset Write and Read Recovery Cycle",
                    "type": "direct",
                    "desc": "Following reset release, immediately initiate wr_en=1 on cycle 1, followed by rd_en=1 on cycle 2.",
                    "exp": "First write is accepted into index 0 (empty deasserts). Read cleanly retrieves written word. Normal FIFO operation fully established."
                }
            ]
        },
        {
            "id": "CAT_BAS",
            "name": "Basic Single & Alternating Access Verification",
            "desc": "Verification of basic single write, single read, data hold characteristics, and alternating write-read sequencing.",
            "tests": [
                {
                    "tc_id": "TC_BAS_001",
                    "name": "Single Write Followed by Single Read",
                    "type": "direct",
                    "desc": "Perform 1 write operation (wr_en=1, rd_en=0, wr_data=65'h1_A5A5_5A5A_A5A5_5A5A). In next cycle, perform 1 read operation (rd_en=1, wr_en=0).",
                    "exp": "On write cycle: empty deasserts (0), full remains 0, occupancy becomes 1. On read cycle: rd_data outputs exact wr_data pattern, empty asserts back to 1."
                },
                {
                    "tc_id": "TC_BAS_002",
                    "name": "Read Data Output Hold During Subsequent Idle Cycles",
                    "type": "direct",
                    "desc": "Write and read a known non-zero word so rd_data displays that word. Maintain wr_en=0 and rd_en=0 for 20 continuous idle clock cycles.",
                    "exp": "rd_data strictly holds its previous accepted read value across all 20 idle cycles. No transition to 0, X, or intermediate state occurs."
                },
                {
                    "tc_id": "TC_BAS_003",
                    "name": "Consecutive Back-to-Back Single Write-Read Pairs (Ping-Pong)",
                    "type": "direct",
                    "desc": "Execute alternating write and read operations for 32 consecutive cycles: [Write D0 -> Read D0 -> Write D1 -> Read D1 ...].",
                    "exp": "Each written word is retrieved in the immediate next read cycle. Occupancy alternates strictly between 0 and 1. empty toggles every cycle; full never asserts."
                },
                {
                    "tc_id": "TC_BAS_004",
                    "name": "Short Burst Write Followed by Short Burst Read",
                    "type": "direct",
                    "desc": "Write 4 consecutive words (occupancy increases 0->1->2->3->4). Then read 4 consecutive words (occupancy decreases 4->3->2->1->0).",
                    "exp": "Burst write smoothly increments occupancy. Burst read retrieves all 4 words in exact FIFO order. empty asserts immediately after the 4th read."
                },
                {
                    "tc_id": "TC_BAS_005",
                    "name": "Zero-Operation Idle Cycles Verification",
                    "type": "direct",
                    "desc": "Apply wr_en=0, rd_en=0 across various occupancy levels (empty, partially full, full).",
                    "exp": "Per Section 3 Operation Table: No operation is executed; rd_data holds; internal occupancy and status flags remain completely unchanged."
                }
            ]
        },
        {
            "id": "CAT_ORD",
            "name": "FIFO Ordering & Data Integrity Verification",
            "desc": "Verification of the fundamental FIFO contract: strict arrival-order preservation, zero data loss, zero duplication, and zero reordering.",
            "tests": [
                {
                    "tc_id": "TC_ORD_001",
                    "name": "Full Depth In-Order Sequencing Verification",
                    "type": "direct",
                    "desc": "Write 64 unique sequential words (0x0001 to 0x0040) to fill the FIFO. Drain and read all 64 words sequentially.",
                    "exp": "All 64 words appear on rd_data in identical arrival order (First-In, First-Out). Zero reordering, zero drops, zero duplications."
                },
                {
                    "tc_id": "TC_ORD_002",
                    "name": "Random Payload Ordering Against Golden Queue",
                    "type": "random",
                    "desc": "Generate 64 randomized 65-bit data payloads. Write all 64 words into FIFO, then read all 64 back, comparing each against a golden queue model.",
                    "exp": "Every read transaction matches the golden software queue entry-for-entry across all 65 data bits with 100% bit-exact accuracy."
                },
                {
                    "tc_id": "TC_ORD_003",
                    "name": "Extended Idle Data Retention & Storage Stability",
                    "type": "direct",
                    "desc": "Fill FIFO with 32 unique words (occupancy=32). Inhibit writes and reads (wr_en=0, rd_en=0) for 500 clock cycles. Then read all 32 words.",
                    "exp": "All 32 items are read back without any data corruption, pointer drift, or memory loss. Flags hold steady throughout the idle period."
                },
                {
                    "tc_id": "TC_ORD_004",
                    "name": "65-Bit Data Bus Full-Width Integrity ([DATA_WIDTH:0])",
                    "type": "direct",
                    "desc": "Write patterns with bit 64 set to 1 (65'h1_0000_0000_0000_0000) and bit 64 set to 0 (65'h0_FFFF_FFFF_FFFF_FFFF). Read back both words.",
                    "exp": "Bit 64 (the 65th declared bit) is fully latched, stored, and retrieved without truncation or bit-slicing errors."
                },
                {
                    "tc_id": "TC_ORD_005",
                    "name": "Variable Latency Read Extraction",
                    "type": "direct",
                    "desc": "Fill FIFO with 16 words. Read out entries with pseudo-random delays (1 to 5 idle cycles between each read).",
                    "exp": "Read data matches FIFO order regardless of read-side latency. rd_data holds valid data during idle wait cycles."
                }
            ]
        },
        {
            "id": "CAT_FUL",
            "name": "FIFO Full & Overflow Handling Verification",
            "desc": "Verification of FIFO filling to maximum capacity (FIFO_DEPTH=64), full flag assertion timing, and overflow write rejection.",
            "tests": [
                {
                    "tc_id": "TC_FUL_001",
                    "name": "Fill to Exact Maximum Capacity (FIFO_DEPTH = 64)",
                    "type": "direct",
                    "desc": "From reset/empty state, drive 64 consecutive write requests (wr_en=1, rd_en=0) with unique incremental data.",
                    "exp": "empty deasserts on 1st write. full asserts precisely on rising clock edge of 64th write. Occupancy reaches 64. rd_data remains 0."
                },
                {
                    "tc_id": "TC_FUL_002",
                    "name": "Almost Full to Full Boundary Transition (Occupancy 63 to 64)",
                    "type": "direct",
                    "desc": "Fill FIFO to 63 entries (FIFO_DEPTH-1). Sample full and empty flags. Apply the 64th write operation.",
                    "exp": "At 63 entries: full=0, empty=0. On 64th write active clock edge: full asserts to 1. Occupancy increments to exactly 64."
                },
                {
                    "tc_id": "TC_FUL_003",
                    "name": "Single Write Overflow Attempt While Full",
                    "type": "direct",
                    "desc": "With FIFO full (full=1), assert wr_en=1 for 1 clock cycle with bogus test data (65'h1_DEAD_BEEF_DEAD_BEEF).",
                    "exp": "Per Section 3: Write is rejected. full remains 1. Internal contents and occupancy remain 64. Bogus data is dropped; rd_data unaffected."
                },
                {
                    "tc_id": "TC_FUL_004",
                    "name": "Sustained Multi-Cycle Overflow Stress",
                    "type": "direct",
                    "desc": "Hold wr_en=1 high for 32 consecutive clock cycles while FIFO remains full (full=1, rd_en=0).",
                    "exp": "All 32 overflow write requests are rejected. full remains solidly high. Write pointer does not increment. Internal memory is uncorrupted."
                },
                {
                    "tc_id": "TC_FUL_005",
                    "name": "Post-Overflow Data Integrity & Drain Verification",
                    "type": "direct",
                    "desc": "Following sustained overflow attempts in TC_FUL_004, read all 64 entries out of the FIFO until empty.",
                    "exp": "All 64 original pre-overflow entries are read out in exact original order. No bogus overflow data appears, and no original data is lost."
                },
                {
                    "tc_id": "TC_FUL_006",
                    "name": "Randomized Overflow Data Stress",
                    "type": "random",
                    "desc": "While FIFO is full, present randomized wr_data patterns with randomized write enable bursts (random wr_en with rd_en=0).",
                    "exp": "No write is accepted. Golden model and DUT remain synchronized at occupancy=64 with full=1."
                }
            ]
        },
        {
            "id": "CAT_EMP",
            "name": "FIFO Empty & Underflow Handling Verification",
            "desc": "Verification of FIFO draining to zero occupancy, empty flag assertion timing, and underflow read rejection.",
            "tests": [
                {
                    "tc_id": "TC_EMP_001",
                    "name": "Drain FIFO from Full to Empty",
                    "type": "direct",
                    "desc": "From full state (64 entries), drive 64 consecutive read requests (rd_en=1, wr_en=0).",
                    "exp": "full deasserts on 1st read edge. empty asserts precisely on 64th read edge. rd_data displays 64th entry. Occupancy becomes 0."
                },
                {
                    "tc_id": "TC_EMP_002",
                    "name": "Almost Empty to Empty Boundary Transition (Occupancy 1 to 0)",
                    "type": "direct",
                    "desc": "Fill FIFO with 1 word (occupancy=1). Verify empty=0, full=0. Apply 1 read operation (rd_en=1).",
                    "exp": "Before edge: empty=0. On active clock edge: rd_data receives the word, occupancy transitions to 0, empty asserts to 1."
                },
                {
                    "tc_id": "TC_EMP_003",
                    "name": "Single Read Underflow Attempt from Reset State",
                    "type": "direct",
                    "desc": "From reset/empty state (empty=1, rd_data=0), assert rd_en=1 for 1 clock cycle.",
                    "exp": "Per Section 3: Read is rejected. rd_data holds 0. empty remains 1. Read pointer does not increment; occupancy stays 0."
                },
                {
                    "tc_id": "TC_EMP_004",
                    "name": "Read Underflow with Prior Valid Data (rd_data Holding Check)",
                    "type": "direct",
                    "desc": "Write a known non-zero word (65'h0_CAFE_BABE_1234_5678). Read it out (empty=1, rd_data=known word). Then assert rd_en=1 while empty.",
                    "exp": "Read is rejected. rd_data strictly continues to hold 65'h0_CAFE_BABE_1234_5678 without changing. empty remains 1."
                },
                {
                    "tc_id": "TC_EMP_005",
                    "name": "Sustained Multi-Cycle Underflow Stress",
                    "type": "direct",
                    "desc": "Hold rd_en=1 high for 32 consecutive clock cycles while FIFO remains empty (empty=1, wr_en=0).",
                    "exp": "All 32 underflow read requests are rejected. empty remains solidly high. rd_data continues holding its previous value without corruption."
                },
                {
                    "tc_id": "TC_EMP_006",
                    "name": "Immediate Write Recovery Following Underflow Stress",
                    "type": "direct",
                    "desc": "Following sustained underflow attempts in TC_EMP_005, assert wr_en=1 with valid data on the very next clock cycle.",
                    "exp": "Write is accepted normally. Occupancy transitions 0 -> 1. empty deasserts on the following active edge. Subsequent read returns this word."
                }
            ]
        },
        {
            "id": "CAT_SRW",
            "name": "Simultaneous Read & Write Operations Verification",
            "desc": "Verification of concurrent read and write operations in partially occupied states, maintaining constant occupancy and strict FIFO ordering.",
            "tests": [
                {
                    "tc_id": "TC_SRW_001",
                    "name": "Simultaneous Read/Write at Mid-Occupancy (Single Cycle)",
                    "type": "direct",
                    "desc": "Pre-fill FIFO to occupancy=16. Assert wr_en=1 and rd_en=1 simultaneously for 1 clock cycle (wr_data=NewData).",
                    "exp": "Per Section 3: Oldest pre-edge item is read to rd_data; NewData is appended to FIFO tail. Occupancy remains exactly 16. empty=0, full=0."
                },
                {
                    "tc_id": "TC_SRW_002",
                    "name": "Continuous Streaming Simultaneous Read/Write (Sustained Steady State)",
                    "type": "direct",
                    "desc": "Pre-fill FIFO to half capacity (32 entries). Drive wr_en=1 and rd_en=1 continuously for 64 consecutive cycles with streaming data.",
                    "exp": "Throughput is exactly 1 word read and 1 word written per cycle. Occupancy remains strictly 32 throughout. All read words match expected sequence."
                },
                {
                    "tc_id": "TC_SRW_003",
                    "name": "Simultaneous Read/Write at Occupancy = 1",
                    "type": "direct",
                    "desc": "FIFO contains exactly 1 entry. Assert wr_en=1 and rd_en=1 simultaneously on rising clock edge.",
                    "exp": "Single stored entry is popped to rd_data; new entry is appended. Occupancy remains 1. empty remains 0, full remains 0."
                },
                {
                    "tc_id": "TC_SRW_004",
                    "name": "Simultaneous Read/Write at Occupancy = FIFO_DEPTH - 1 (63)",
                    "type": "direct",
                    "desc": "FIFO contains exactly 63 entries. Assert wr_en=1 and rd_en=1 simultaneously on rising clock edge.",
                    "exp": "Oldest entry is popped to rd_data; new entry is appended. Occupancy remains 63. empty remains 0, full remains 0."
                },
                {
                    "tc_id": "TC_SRW_005",
                    "name": "Random Simultaneous Read/Write Across Arbitrary Occupancies",
                    "type": "random",
                    "desc": "At arbitrary random occupancy levels (between 2 and 62), inject random simultaneous read/write bursts.",
                    "exp": "For every simultaneous cycle, read value matches pre-edge oldest entry, write is accepted, and occupancy remains constant."
                }
            ]
        },
        {
            "id": "CAT_BND",
            "name": "Boundary Transitions & Corner Case Verification",
            "desc": "Verification of critical boundary corner cases explicitly specified in Section 3 and Section 6: simultaneous operations when empty or full.",
            "tests": [
                {
                    "tc_id": "TC_BND_001",
                    "name": "Simultaneous Read/Write on EMPTY FIFO Boundary (empty=1, full=0)",
                    "type": "direct",
                    "desc": "With FIFO completely empty (empty=1, full=0), assert both wr_en=1 and rd_en=1 simultaneously on the rising clock edge.",
                    "exp": "Per Section 3 Table: Read is rejected (rd_data holds previous value). Write is accepted (full=0). Occupancy transitions 0->1; empty deasserts next cycle."
                },
                {
                    "tc_id": "TC_BND_002",
                    "name": "Simultaneous Read/Write on FULL FIFO Boundary (full=1, empty=0)",
                    "type": "direct",
                    "desc": "With FIFO completely full (full=1, empty=0), assert both wr_en=1 and rd_en=1 simultaneously on the rising clock edge.",
                    "exp": "Per Section 3 Table: Write is rejected (full=1, incoming data dropped). Read is accepted (empty=0, oldest item to rd_data). Occupancy transitions 64->63; full deasserts next cycle."
                },
                {
                    "tc_id": "TC_BND_003",
                    "name": "Empty Boundary Thrash (Rapid 0 <-> 1 Occupancy Oscillation)",
                    "type": "direct",
                    "desc": "Oscillate occupancy around empty boundary for 30 cycles: [Write 1 -> Read 1 -> Write 1 -> Read 1 ...].",
                    "exp": "empty flag toggles between 1 and 0 every cycle without glitching or desynchronization. FIFO never enters underflow or overflow."
                },
                {
                    "tc_id": "TC_BND_004",
                    "name": "Full Boundary Thrash (Rapid 63 <-> 64 Occupancy Oscillation)",
                    "type": "direct",
                    "desc": "Fill FIFO to 63 entries, then alternate: [Write to 64 -> Read to 63 -> Write to 64 -> Read to 63 ...] for 30 cycles.",
                    "exp": "full flag toggles between 1 and 0 every cycle cleanly. No write is lost or corrupted; all reads match expected order."
                },
                {
                    "tc_id": "TC_BND_005",
                    "name": "Simultaneous Read/Write Boundary State Recovery",
                    "type": "direct",
                    "desc": "Immediately after empty boundary simultaneous operation (TC_BND_001), verify the newly written entry can be read out cleanly.",
                    "exp": "The item written during the empty-boundary cycle is read out with 100% data integrity; rd_data updates and empty re-asserts."
                },
                {
                    "tc_id": "TC_BND_006",
                    "name": "Full Boundary Rejected Write Non-Pollution Verification",
                    "type": "direct",
                    "desc": "Following full boundary simultaneous operation (TC_BND_002), drain all 63 remaining items and verify the rejected write was not appended.",
                    "exp": "Exactly 63 pre-existing items are retrieved. The rejected write word does not appear anywhere in the FIFO stream."
                }
            ]
        },
        {
            "id": "CAT_DAT",
            "name": "Data Patterns & Width Stress Verification",
            "desc": "Verification of data pattern corner cases specified in Section 4: all zeros, all ones, alternating bits, walking 1s, walking 0s, and 65-bit bus coverage.",
            "tests": [
                {
                    "tc_id": "TC_DAT_001",
                    "name": "All-Zeros Data Pattern Across Full 65-Bit Width",
                    "type": "direct",
                    "desc": "Write wr_data = 65'h0_0000_0000_0000_0000 into FIFO. Read it back and verify rd_data.",
                    "exp": "All 65 bits are written and retrieved as zeros. rd_data updates to 0 and empty flag updates properly."
                },
                {
                    "tc_id": "TC_DAT_002",
                    "name": "All-Ones Data Pattern Across Full 65-Bit Width",
                    "type": "direct",
                    "desc": "Write wr_data = 65'h1_FFFF_FFFF_FFFF_FFFF into FIFO. Read it back and verify rd_data.",
                    "exp": "All 65 bits are written and retrieved as ones. Every bit from 0 to 64 is verified high without stuck-at-zero faults."
                },
                {
                    "tc_id": "TC_DAT_003",
                    "name": "Alternating Bit Patterns (Checkerboard 0xAA / 0x55)",
                    "type": "direct",
                    "desc": "Write word 1: 65'h0_AAAA_AAAA_AAAA_AAAA; Write word 2: 65'h1_5555_5555_5555_5555. Read back both words.",
                    "exp": "Both checkerboard patterns are read back intact, verifying no crosstalk, adjacent bit shorts, or coupling between bits."
                },
                {
                    "tc_id": "TC_DAT_004",
                    "name": "Walking-Ones Pattern Across All 65 Bits",
                    "type": "direct",
                    "desc": "Sequentially write 65 individual words where a single '1' walks from bit 0 up to bit 64 (wr_data = 1 << i, for i=0..64). Read all back.",
                    "exp": "All 65 words are read back matching exactly, verifying that every single bit line (0 through 64) is individually controllable and observable."
                },
                {
                    "tc_id": "TC_DAT_005",
                    "name": "Walking-Zeros Pattern Across All 65 Bits",
                    "type": "direct",
                    "desc": "Sequentially write 65 individual words where a single '0' walks from bit 0 up to bit 64 (wr_data = ~(1 << i), for i=0..64). Read all back.",
                    "exp": "All 65 words are read back matching exactly, verifying that no bit line is stuck at 1 across the full 65-bit bus."
                },
                {
                    "tc_id": "TC_DAT_006",
                    "name": "Pseudo-Random Data Distribution Stress",
                    "type": "random",
                    "desc": "Write and read a sequence of 64 randomized 65-bit numbers with random bit densities (sparse 1s, dense 1s, random hashes).",
                    "exp": "All random patterns match golden model values bit-for-bit upon read."
                }
            ]
        },
        {
            "id": "CAT_RND",
            "name": "Randomized Traffic & Stress Verification",
            "desc": "Constrained random verification, heavy traffic distributions, extended stability runs, and continuous scoreboard checking against golden queue.",
            "tests": [
                {
                    "tc_id": "TC_RND_001",
                    "name": "Balanced Randomized Traffic (P(wr)=0.5, P(rd)=0.5)",
                    "type": "random",
                    "desc": "Run 2,500 clock cycles of randomized traffic with equal write and read probabilities. Random payload data on each write.",
                    "exp": "Scoreboard continuously checks all accepted reads against golden reference queue. Occupancy randomly traverses 0 to 64. Zero mismatches."
                },
                {
                    "tc_id": "TC_RND_002",
                    "name": "Write-Heavy Randomized Traffic (P(wr)=0.85, P(rd)=0.15)",
                    "type": "random",
                    "desc": "Run 2,000 clock cycles with 85% write probability to repeatedly saturate FIFO, stressing full flag generation and overflow rejection.",
                    "exp": "full flag consistently prevents writes when full; all accepted items match golden model; rejected writes do not corrupt storage."
                },
                {
                    "tc_id": "TC_RND_003",
                    "name": "Read-Heavy Randomized Traffic (P(wr)=0.15, P(rd)=0.85)",
                    "type": "random",
                    "desc": "Run 2,000 clock cycles with 85% read probability to repeatedly drain FIFO, stressing empty flag generation and underflow rejection.",
                    "exp": "empty flag consistently prevents reads when empty; rd_data holds previous value on rejected reads; zero scoreboard errors."
                },
                {
                    "tc_id": "TC_RND_004",
                    "name": "Random Bursty Producer / Consumer Traffic",
                    "type": "random",
                    "desc": "Random write bursts of length 1..64 followed by random idle delays, interleaved with random read bursts of length 1..64.",
                    "exp": "Burst boundaries transition seamlessly; flag timing adheres to spec; queue occupancy and data integrity match golden reference."
                },
                {
                    "tc_id": "TC_RND_005",
                    "name": "Long-Duration Extended Stress & Pointer Wraparound Test",
                    "type": "random",
                    "desc": "Run 10,000+ continuous randomized transactions to force internal read and write pointers to wrap around address space >150 times.",
                    "exp": "Zero memory leaks, zero pointer drift or desynchronization. FIFO maintains 100% data integrity and correct flag status across all wraparounds."
                },
                {
                    "tc_id": "TC_RND_006",
                    "name": "Randomized Traffic with Asynchronous Reset Injection",
                    "type": "random",
                    "desc": "During active randomized write/read traffic, assert asynchronous reset (sys_rst=0) at random time points. Release reset and resume traffic.",
                    "exp": "FIFO immediately purges all state upon reset; reference model clears queue; post-reset operations restart cleanly from empty state."
                }
            ]
        },
        {
            "id": "CAT_PRM",
            "name": "Parameter Configuration & Scalability Verification",
            "desc": "Verification of design scalability across parameter configurations noted in Section 1 and Section 6 (differing depths, widths, and power-of-two behavior).",
            "tests": [
                {
                    "tc_id": "TC_PRM_001",
                    "name": "Minimum Depth FIFO Corner Case (FIFO_DEPTH = 2)",
                    "type": "direct",
                    "desc": "Instantiate DUT with FIFO_DEPTH=2. Test: write 1, write 2 (full asserts), read 1, read 2 (empty asserts), and simultaneous read/write.",
                    "exp": "full asserts on exactly the 2nd write. Simultaneous read/write keeps occupancy at 1. Boundary operations function flawlessly."
                },
                {
                    "tc_id": "TC_PRM_002",
                    "name": "Small Power-of-Two Depth Configuration (FIFO_DEPTH = 4)",
                    "type": "direct",
                    "desc": "Instantiate DUT with FIFO_DEPTH=4. Execute full fill, overflow rejection, full drain, underflow rejection, and pointer wraparound.",
                    "exp": "Pointer wrapping occurs at modulo 4; full and empty flags reflect depth 4 accurately."
                },
                {
                    "tc_id": "TC_PRM_003",
                    "name": "Non-Power-of-Two Depth Wraparound Test (FIFO_DEPTH = 5 or 10)",
                    "type": "direct",
                    "desc": "Instantiate DUT with non-power-of-two depth (e.g., FIFO_DEPTH=5) as highlighted in Section 6 open points. Fill to 5, overflow, and drain.",
                    "exp": "Confirm pointer wraparound logic correctly rolls over at count 5 without illegal states or premature flag assertions."
                },
                {
                    "tc_id": "TC_PRM_004",
                    "name": "Narrow Data Width Parameter Configuration (DATA_WIDTH = 8, 9-bit bus)",
                    "type": "direct",
                    "desc": "Instantiate DUT with DATA_WIDTH=8 (declared port [8:0], 9 bits wide). Verify all 9 bits with walking ones and alternating patterns.",
                    "exp": "9-bit data paths store and retrieve data with zero upper-bit leakage or truncation."
                },
                {
                    "tc_id": "TC_PRM_005",
                    "name": "Wide Data Width Parameter Configuration (DATA_WIDTH = 128, 129-bit bus)",
                    "type": "direct",
                    "desc": "Instantiate DUT with DATA_WIDTH=128 (declared port [128:0], 129 bits wide). Verify full 129-bit transactions and boundary flags.",
                    "exp": "129-bit data path operates without truncation; bit 128 (MSB) latched and retrieved accurately."
                }
            ]
        }
    ]

    # Calculate summary metrics
    total_tests = sum(len(cat["tests"]) for cat in categories)
    direct_tests = sum(sum(1 for t in cat["tests"] if t["type"] == "direct") for cat in categories)
    random_tests = sum(sum(1 for t in cat["tests"] if t["type"] == "random") for cat in categories)

    # ==========================================
    # SHEET 1: OVERVIEW DASHBOARD
    # ==========================================
    ws_dash = wb.create_sheet(title="Overview & Summary")
    ws_dash.views.sheetView[0].showGridLines = True
    
    # Title
    ws_dash["B2"] = "Sync_FIFO_SBRX - Verification Plan & Test Specification"
    ws_dash["B2"].font = DASH_TITLE_FONT
    ws_dash["B3"] = "Black-Box Verification Plan according to Sync_FIFO_SBRX Specification (Clear Edition)"
    ws_dash["B3"].font = Font(name="Calibri", size=11, italic=True, color="555555")

    # Metadata Card
    ws_dash["B5"] = "1. DESIGN SPECIFICATION & INTERFACE SUMMARY"
    ws_dash["B5"].font = SECTION_TITLE_FONT
    
    spec_meta = [
        ("Block Type", "Single-clock synchronous FIFO (Sync_FIFO_SBRX)"),
        ("Clock Domain", "lclk; all operations evaluated on rising clock edge"),
        ("Reset Type", "sys_rst; active-low asynchronous reset (sys_rst=0 resets FIFO)"),
        ("Default DATA_WIDTH", "64 (Declared ports [DATA_WIDTH:0] -> 65 bits wide)"),
        ("Default FIFO_DEPTH", "64 entries (Deep memory buffer)"),
        ("Status Flags", "empty (high when 0 entries), full (high when FIFO_DEPTH entries)"),
        ("Output Port rd_data", "Most recent accepted read value; holds during idle/underflow; resets to 0"),
        ("Contract & Gating", "Independent gating: writes gated by full, reads gated by empty"),
        ("Sign-Off Criterion", "All directed and randomized tests pass; zero data loss/duplication/reorder")
    ]
    
    for idx, (k, v) in enumerate(spec_meta, start=6):
        c1 = ws_dash.cell(row=idx, column=2, value=k)
        c2 = ws_dash.cell(row=idx, column=3, value=v)
        c1.font = BOLD_FONT
        c1.fill = ALT_ROW_FILL
        c1.border = THIN_BORDER
        c2.font = REGULAR_FONT
        c2.border = THIN_BORDER
        ws_dash.merge_cells(start_row=idx, start_column=3, end_row=idx, end_column=6)

    # Metrics Summary Table
    ws_dash["B17"] = "2. VERIFICATION METRICS & CATEGORY BREAKDOWN"
    ws_dash["B17"].font = SECTION_TITLE_FONT
    
    metrics_headers = ["Category ID", "Category Name", "Direct Tests", "Random Tests", "Total Tests"]
    for col_idx, h in enumerate(metrics_headers, start=2):
        cell = ws_dash.cell(row=18, column=col_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = ALIGN_CENTER
        cell.border = HEADER_BORDER
    
    curr_row = 19
    for cat in categories:
        c_direct = sum(1 for t in cat["tests"] if t["type"] == "direct")
        c_random = sum(1 for t in cat["tests"] if t["type"] == "random")
        c_total = len(cat["tests"])
        
        row_cells = [
            ws_dash.cell(row=curr_row, column=2, value=cat["id"]),
            ws_dash.cell(row=curr_row, column=3, value=cat["name"]),
            ws_dash.cell(row=curr_row, column=4, value=c_direct),
            ws_dash.cell(row=curr_row, column=5, value=c_random),
            ws_dash.cell(row=curr_row, column=6, value=c_total)
        ]
        row_cells[0].alignment = ALIGN_CENTER
        row_cells[0].font = CODE_FONT
        row_cells[1].alignment = Alignment(horizontal="left", vertical="center")
        row_cells[1].font = REGULAR_FONT
        row_cells[2].alignment = ALIGN_CENTER
        row_cells[2].font = REGULAR_FONT
        row_cells[3].alignment = ALIGN_CENTER
        row_cells[3].font = REGULAR_FONT
        row_cells[4].alignment = ALIGN_CENTER
        row_cells[4].font = BOLD_FONT
        
        fill = ALT_ROW_FILL if curr_row % 2 == 0 else WHITE_ROW_FILL
        for c in row_cells:
            c.border = THIN_BORDER
            if c != row_cells[0]:
                c.fill = fill
        curr_row += 1
        
    # Total Row
    tot_cells = [
        ws_dash.cell(row=curr_row, column=2, value="TOTAL"),
        ws_dash.cell(row=curr_row, column=3, value="All 10 Categories Combined"),
        ws_dash.cell(row=curr_row, column=4, value=direct_tests),
        ws_dash.cell(row=curr_row, column=5, value=random_tests),
        ws_dash.cell(row=curr_row, column=6, value=total_tests)
    ]
    for c in tot_cells:
        c.font = HEADER_FONT
        c.fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
        c.alignment = ALIGN_CENTER
        c.border = HEADER_BORDER
    tot_cells[1].alignment = Alignment(horizontal="left", vertical="center")

    # Core Reference Model Checks section
    curr_row += 2
    ws_dash.cell(row=curr_row, column=2, value="3. GOLDEN REFERENCE MODEL & CORE CHECKING RULES (SPEC SECTION 5)").font = SECTION_TITLE_FONT
    curr_row += 1
    
    check_rules = [
        ("Rule 1 - Empty Flag", "empty = 1 strictly whenever golden reference queue occupancy is zero; empty = 0 otherwise."),
        ("Rule 2 - Full Flag", "full = 1 strictly whenever golden reference queue occupancy equals FIFO_DEPTH (64); full = 0 otherwise."),
        ("Rule 3 - Occupancy Bounds", "Reference queue occupancy strictly stays bounded within [0, FIFO_DEPTH] at all active edges."),
        ("Rule 4 - In-Order Delivery", "Every accepted read (rd_en=1 & empty=0) returns the exact oldest queued item in FIFO order."),
        ("Rule 5 - Read Rejection Hold", "A rejected read (rd_en=1 & empty=1) does NOT change rd_data; rd_data strictly holds previous value."),
        ("Rule 6 - Write Rejection", "A rejected write (wr_en=1 & full=1) does NOT alter the stored sequence or internal pointers."),
        ("Rule 7 - Occupancy Conservation", "Accepted writes minus accepted reads strictly equals current occupancy since reset assertion.")
    ]
    for r_title, r_desc in check_rules:
        c1 = ws_dash.cell(row=curr_row, column=2, value=r_title)
        c2 = ws_dash.cell(row=curr_row, column=3, value=r_desc)
        c1.font = BOLD_FONT
        c1.fill = ALT_ROW_FILL
        c1.border = THIN_BORDER
        c2.font = REGULAR_FONT
        c2.border = THIN_BORDER
        ws_dash.merge_cells(start_row=curr_row, start_column=3, end_row=curr_row, end_column=6)
        curr_row += 1

    # Set column widths for Dashboard
    ws_dash.column_dimensions["A"].width = 4
    ws_dash.column_dimensions["B"].width = 24
    ws_dash.column_dimensions["C"].width = 42
    ws_dash.column_dimensions["D"].width = 16
    ws_dash.column_dimensions["E"].width = 16
    ws_dash.column_dimensions["F"].width = 16

    # ==========================================
    # SHEET 2: VERIFICATION PLAN (MASTER SHEET WITH EXACT 5 HEADERS)
    # Headers exactly matching user's image:
    # [TC ID] [Test Name] [Test Type] [Description] [Expected Result]
    # Split into categories with prominent Category Banner Rows
    # ==========================================
    ws_plan = wb.create_sheet(title="Verification_Plan")
    ws_plan.views.sheetView[0].showGridLines = True
    
    headers_5 = ["TC ID", "Test Name", "Test Type", "Description", "Expected Result"]
    
    # Table header at row 1
    ws_plan.row_dimensions[1].height = 28
    for col_idx, h in enumerate(headers_5, start=1):
        cell = ws_plan.cell(row=1, column=col_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = ALIGN_CENTER
        cell.border = HEADER_BORDER
        
    p_row = 2
    for cat_idx, cat in enumerate(categories, start=1):
        # Category Banner Row spanning cols A:E
        ws_plan.row_dimensions[p_row].height = 25
        cat_banner_text = f"CATEGORY {cat_idx}: {cat['name'].upper()}  -  {cat['desc']}"
        banner_cell = ws_plan.cell(row=p_row, column=1, value=cat_banner_text)
        banner_cell.font = CAT_FONT
        banner_cell.fill = CAT_FILL
        banner_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for c_idx in range(1, 6):
            c = ws_plan.cell(row=p_row, column=c_idx)
            c.fill = CAT_FILL
            c.border = THIN_BORDER
        ws_plan.merge_cells(start_row=p_row, start_column=1, end_row=p_row, end_column=5)
        p_row += 1
        
        # Test Cases in this Category
        for t_idx, test in enumerate(cat["tests"]):
            ws_plan.row_dimensions[p_row].height = 42
            
            c_id = ws_plan.cell(row=p_row, column=1, value=test["tc_id"])
            c_name = ws_plan.cell(row=p_row, column=2, value=test["name"])
            c_type = ws_plan.cell(row=p_row, column=3, value=test["type"])
            c_desc = ws_plan.cell(row=p_row, column=4, value=test["desc"])
            c_exp = ws_plan.cell(row=p_row, column=5, value=test["exp"])
            
            # Styling
            c_id.alignment = ALIGN_CENTER
            c_id.font = CODE_FONT
            
            c_name.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            c_name.font = BOLD_FONT
            
            c_type.alignment = ALIGN_CENTER
            c_type.font = BOLD_FONT
            if test["type"].lower() == "direct":
                c_type.fill = DIRECT_FILL
            else:
                c_type.fill = RANDOM_FILL
                
            c_desc.alignment = ALIGN_LEFT
            c_desc.font = REGULAR_FONT
            
            c_exp.alignment = ALIGN_LEFT
            c_exp.font = REGULAR_FONT
            
            row_fill = ALT_ROW_FILL if t_idx % 2 == 1 else WHITE_ROW_FILL
            for cell in [c_id, c_name, c_desc, c_exp]:
                cell.border = THIN_BORDER
                if cell != c_type:
                    cell.fill = row_fill
            c_type.border = THIN_BORDER
            
            p_row += 1

    # Column widths for Verification_Plan
    ws_plan.column_dimensions["A"].width = 16
    ws_plan.column_dimensions["B"].width = 34
    ws_plan.column_dimensions["C"].width = 15
    ws_plan.column_dimensions["D"].width = 54
    ws_plan.column_dimensions["E"].width = 54

    # Enable freeze panes on Verification_Plan
    ws_plan.freeze_panes = "A2"

    # ==========================================
    # SHEET 3: COMPLETE FILTERABLE TEST MATRIX (WITH CATEGORY COLUMN)
    # Headers: [Category] [TC ID] [Test Name] [Test Type] [Description] [Expected Result]
    # Provides explicit filtering capability by Category, Test Type, etc.
    # ==========================================
    ws_matrix = wb.create_sheet(title="Filterable_Matrix")
    ws_matrix.views.sheetView[0].showGridLines = True
    
    matrix_headers = ["Category", "TC ID", "Test Name", "Test Type", "Description", "Expected Result"]
    ws_matrix.row_dimensions[1].height = 28
    for col_idx, h in enumerate(matrix_headers, start=1):
        cell = ws_matrix.cell(row=1, column=col_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = ALIGN_CENTER
        cell.border = HEADER_BORDER
        
    m_row = 2
    for cat in categories:
        for t_idx, test in enumerate(cat["tests"]):
            ws_matrix.row_dimensions[m_row].height = 38
            
            c_cat = ws_matrix.cell(row=m_row, column=1, value=cat["name"])
            c_id = ws_matrix.cell(row=m_row, column=2, value=test["tc_id"])
            c_name = ws_matrix.cell(row=m_row, column=3, value=test["name"])
            c_type = ws_matrix.cell(row=m_row, column=4, value=test["type"])
            c_desc = ws_matrix.cell(row=m_row, column=5, value=test["desc"])
            c_exp = ws_matrix.cell(row=m_row, column=6, value=test["exp"])
            
            c_cat.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            c_cat.font = BOLD_FONT
            
            c_id.alignment = ALIGN_CENTER
            c_id.font = CODE_FONT
            
            c_name.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            c_name.font = REGULAR_FONT
            
            c_type.alignment = ALIGN_CENTER
            c_type.font = BOLD_FONT
            if test["type"].lower() == "direct":
                c_type.fill = DIRECT_FILL
            else:
                c_type.fill = RANDOM_FILL
                
            c_desc.alignment = ALIGN_LEFT
            c_desc.font = REGULAR_FONT
            
            c_exp.alignment = ALIGN_LEFT
            c_exp.font = REGULAR_FONT
            
            row_fill = ALT_ROW_FILL if m_row % 2 == 1 else WHITE_ROW_FILL
            for cell in [c_cat, c_id, c_name, c_desc, c_exp]:
                cell.border = THIN_BORDER
                cell.fill = row_fill
            c_type.border = THIN_BORDER
            
            m_row += 1
            
    # Auto-filter for Matrix
    ws_matrix.auto_filter.ref = f"A1:F{m_row-1}"
    ws_matrix.freeze_panes = "A2"
    
    ws_matrix.column_dimensions["A"].width = 28
    ws_matrix.column_dimensions["B"].width = 16
    ws_matrix.column_dimensions["C"].width = 34
    ws_matrix.column_dimensions["D"].width = 15
    ws_matrix.column_dimensions["E"].width = 54
    ws_matrix.column_dimensions["F"].width = 54

    # ==========================================
    # DEDICATED INDIVIDUAL CATEGORY SHEETS
    # Each category gets its own clean tab with the exact 5 headers
    # ==========================================
    tab_short_names = [
        "1_Reset_Init",
        "2_Basic_Access",
        "3_Ordering_Integrity",
        "4_Capacity_Full",
        "5_Drain_Empty",
        "6_Simultaneous_RW",
        "7_Boundary_Corners",
        "8_Data_Patterns",
        "9_Random_Stress",
        "10_Param_Scalability"
    ]
    
    for cat_idx, cat in enumerate(categories):
        tab_name = tab_short_names[cat_idx]
        ws_cat = wb.create_sheet(title=tab_name)
        ws_cat.views.sheetView[0].showGridLines = True
        
        # Category Title Banner at row 1
        ws_cat.row_dimensions[1].height = 26
        ws_cat.cell(row=1, column=1, value=f"{cat['name'].upper()}").font = Font(name="Calibri", size=13, bold=True, color="1B4F72")
        ws_cat.row_dimensions[2].height = 18
        ws_cat.cell(row=2, column=1, value=f"Objective: {cat['desc']}").font = Font(name="Calibri", size=10, italic=True, color="555555")
        
        # Headers at row 4
        ws_cat.row_dimensions[4].height = 28
        for col_idx, h in enumerate(headers_5, start=1):
            cell = ws_cat.cell(row=4, column=col_idx, value=h)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = ALIGN_CENTER
            cell.border = HEADER_BORDER
            
        c_row = 5
        for t_idx, test in enumerate(cat["tests"]):
            ws_cat.row_dimensions[c_row].height = 42
            
            c_id = ws_cat.cell(row=c_row, column=1, value=test["tc_id"])
            c_name = ws_cat.cell(row=c_row, column=2, value=test["name"])
            c_type = ws_cat.cell(row=c_row, column=3, value=test["type"])
            c_desc = ws_cat.cell(row=c_row, column=4, value=test["desc"])
            c_exp = ws_cat.cell(row=c_row, column=5, value=test["exp"])
            
            c_id.alignment = ALIGN_CENTER
            c_id.font = CODE_FONT
            
            c_name.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            c_name.font = BOLD_FONT
            
            c_type.alignment = ALIGN_CENTER
            c_type.font = BOLD_FONT
            if test["type"].lower() == "direct":
                c_type.fill = DIRECT_FILL
            else:
                c_type.fill = RANDOM_FILL
                
            c_desc.alignment = ALIGN_LEFT
            c_desc.font = REGULAR_FONT
            
            c_exp.alignment = ALIGN_LEFT
            c_exp.font = REGULAR_FONT
            
            fill = ALT_ROW_FILL if t_idx % 2 == 1 else WHITE_ROW_FILL
            for cell in [c_id, c_name, c_desc, c_exp]:
                cell.border = THIN_BORDER
                cell.fill = fill
            c_type.border = THIN_BORDER
            
            c_row += 1
            
        ws_cat.auto_filter.ref = f"A4:E{c_row-1}"
        ws_cat.freeze_panes = "A5"
        
        ws_cat.column_dimensions["A"].width = 16
        ws_cat.column_dimensions["B"].width = 34
        ws_cat.column_dimensions["C"].width = 15
        ws_cat.column_dimensions["D"].width = 54
        ws_cat.column_dimensions["E"].width = 54

    # Save workbook
    output_path = r"d:\ASU\Senior 2\Graduation Project\IPs\Sync_fifo_SBRX\Sync_FIFO_SBRX_Verification_Plan.xlsx"
    wb.save(output_path)
    print(f"Successfully generated verification plan: {output_path}")
    print(f"Total Test Cases: {total_tests} (Direct: {direct_tests}, Random: {random_tests}) across {len(categories)} categories.")

if __name__ == "__main__":
    build_verification_plan()
